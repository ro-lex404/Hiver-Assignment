import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def clean_latex(text: str) -> str:
    text = text.replace(r"$\rightarrow$", "→")
    text = text.replace(r"$ightarrow$", "→")
    text = text.replace(r"ightarrow", "→")
    text = text.replace(r"\rightarrow", "→")
    text = text.replace(r"$\le$", "≤")
    text = text.replace(r"\le", "≤")
    text = text.replace(r"$\ge$", "≥")
    text = text.replace(r"\ge", "≥")
    text = text.replace(r"$\times$", "×")
    text = text.replace(r"\times", "×")
    text = text.replace(r"$\alpha=32$", "α=32")
    text = text.replace(r"$\alpha$", "α")
    text = text.replace(r"\alpha", "α")
    text = text.replace(r"$\kappa$", "κ")
    text = text.replace(r"\kappa", "κ")
    text = text.replace(r"$\approx$", "≈")
    text = text.replace(r"\approx", "≈")
    text = text.replace(r"$\text{risk} < 0.65$", "(risk < 0.65)")
    text = text.replace(r"$\text{risk} \ge 0.65$", "(risk ≥ 0.65)")
    text = text.replace(r"$\text{risk}$", "risk")
    text = text.replace(r"\text{risk}", "risk")
    text = text.replace(r"$\text{confidence} < 0.70 \rightarrow +0.50$", "(confidence < 0.70 → +0.50)")
    text = text.replace(r"\text{confidence}", "confidence")
    text = text.replace(r"$r=16$", "r=16")
    text = text.replace(r"$2\times 10^{-4}$", "2 × 10⁻⁴")
    text = text.replace(r"$0.50 \le \text{confidence} < 0.70$", "(0.50 ≤ confidence < 0.70)")
    text = text.replace(r"$\le 280$", "≤ 280")
    text = text.replace(r"\le 280", "≤ 280")
    text = re.sub(r"\$([^\$]+)\$", r"\1", text)
    return text

def add_styled_paragraph(doc, text, style='Normal', space_after=6, space_before=0, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    _apply_inline_formatting(p, text)
    return p

def _apply_inline_formatting(paragraph, text):
    text = clean_latex(text)
    pattern = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|https?://[^\s)]+)')
    tokens = pattern.split(text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('*') and token.endswith('*'):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(180, 40, 40)
        elif token.startswith('http://') or token.startswith('https://'):
            run = paragraph.add_run(token)
            run.font.color.rgb = RGBColor(31, 78, 121)
            run.underline = True
        else:
            paragraph.add_run(token)

def convert_md_to_docx(md_path: str, docx_path: str):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(38, 38, 38)

    in_code_block = False
    code_lines = []
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal table_rows, in_table
        if not table_rows:
            in_table = False
            return
        header = table_rows[0]
        data_rows = table_rows[1:]
        cols = len(header)
        tbl = doc.add_table(rows=len(table_rows), cols=cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = True
        for col_idx, text in enumerate(header):
            cell = tbl.cell(0, col_idx)
            set_cell_background(cell, '1F4E79')
            set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text.strip())
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(10)
        for r_idx, row_data in enumerate(data_rows):
            bg_color = 'F9FAFB' if r_idx % 2 == 1 else 'FFFFFF'
            for col_idx in range(cols):
                cell = tbl.cell(r_idx + 1, col_idx)
                set_cell_background(cell, bg_color)
                set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                cell_text = row_data[col_idx].strip() if col_idx < len(row_data) else ''
                _apply_inline_formatting(p, cell_text)
                for run in p.runs:
                    run.font.size = Pt(9.5)
        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_after = Pt(6)
        table_rows = []
        in_table = False

    def flush_code():
        nonlocal code_lines, in_code_block
        if not code_lines:
            in_code_block = False
            return
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, 'F2F4F7')
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run('\n'.join(code_lines))
        run.font.name = 'Consolas'
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(30, 30, 30)
        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_after = Pt(4)
        code_lines = []
        in_code_block = False

    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\r\n')
        stripped = line.strip()
        if stripped.startswith('```'):
            if in_code_block:
                flush_code()
            else:
                if in_table:
                    flush_table()
                in_code_block = True
                code_lines = []
            i += 1
            continue
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue
        if '|' in stripped and stripped.startswith('|') and stripped.endswith('|'):
            if re.match(r'^\|(?:\s*:?-+:?\s*\|)+$', stripped):
                i += 1
                continue
            cells = [c.strip() for c in stripped.strip('|').split('|')]
            table_rows.append(cells)
            in_table = True
            i += 1
            continue
        elif in_table:
            flush_table()
        if not stripped:
            i += 1
            continue
        if stripped in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="D3D3D3"/></w:pBdr>')
            p._p.get_or_add_pPr().append(pBdr)
            i += 1
            continue
        if stripped.startswith('> '):
            quote_text = stripped[2:].strip()
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.cell(0, 0)
            set_cell_background(cell, 'EDF2F7')
            set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            _apply_inline_formatting(p, quote_text)
            for run in p.runs:
                run.italic = True
                run.font.size = Pt(10)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            i += 1
            continue
        if stripped.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(clean_latex(stripped[2:].strip()))
            run.font.name = 'Calibri Light'
            run.font.size = Pt(22)
            run.bold = True
            run.font.color.rgb = RGBColor(31, 78, 121)
            i += 1
            continue
        elif stripped.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(clean_latex(stripped[3:].strip()))
            run.font.name = 'Calibri Light'
            run.font.size = Pt(15)
            run.bold = True
            run.font.color.rgb = RGBColor(31, 78, 121)
            i += 1
            continue
        elif stripped.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(clean_latex(stripped[4:].strip()))
            run.font.name = 'Calibri'
            run.font.size = Pt(12)
            run.bold = True
            run.font.color.rgb = RGBColor(68, 114, 196)
            i += 1
            continue
        elif stripped.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(clean_latex(stripped[5:].strip()))
            run.font.name = 'Calibri'
            run.font.size = Pt(11)
            run.bold = True
            run.font.color.rgb = RGBColor(89, 89, 89)
            i += 1
            continue
        if stripped.startswith('- ') or stripped.startswith('* ') or stripped.startswith('• '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            _apply_inline_formatting(p, stripped[2:].strip())
            i += 1
            continue
        m_num = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if m_num:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            _apply_inline_formatting(p, m_num.group(2).strip())
            i += 1
            continue
        add_styled_paragraph(doc, stripped, space_before=0, space_after=5)
        i += 1

    if in_table:
        flush_table()
    if in_code_block:
        flush_code()

    doc.save(docx_path)
    print(f'Successfully generated {docx_path}')

if __name__ == '__main__':
    convert_md_to_docx('reports/REPORT.md', 'reports/REPORT.docx')
